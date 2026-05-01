# app/api/export.py

from fastapi import APIRouter, HTTPException
from supabase import create_client
import os
from datetime import datetime
from fastapi.responses import StreamingResponse


from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether

from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
)
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
import io

import zipfile
import json


from pptx import Presentation
from pptx.util import Inches

router = APIRouter()

supabase = create_client(
    os.getenv("SUPABASE_URL"),
    os.getenv("SUPABASE_KEY")
)


@router.get("/api/exports/{project_id}")
def export_project(project_id: str):

    # --- FETCH PHASE RUNS ---
    phase_res = (
        supabase.table("phase_runs")
        .select("*")
        .eq("project_id", project_id)
        .eq("status", "completed")
        .order("phase_number", desc=False)
        .execute()
    )

    if not phase_res.data:
        raise HTTPException(404, "No phase data found")

    # --- FETCH BOARDROOM RUNS ---
    boardroom_res = (
        supabase.table("boardroom_runs")
        .select("*")
        .eq("project_id", project_id)
        .order("created_at", desc=False)
        .execute()
    )

    # --- TRANSFORM ---
    phases = []
    for row in phase_res.data:
        phases.append({
            "phase_number": row["phase_number"],
            "retry_number": row["retry_number"],
            "output": row["api_output"] or row["raw_output"]
        })

    boardroom = []
    for row in boardroom_res.data:
        boardroom.append({
            "phase": row["phase_number"],
            "status": row["status"],
            "output": row.get("api_output") or row.get("raw_output")
        })

    # --- FINAL MANIFEST ---
    manifest = {
        "project_id": project_id,
        "generated_at": datetime.utcnow().isoformat(),
        "phases": phases,
        "boardroom": boardroom
    }

    return manifest


@router.get("/api/exports/{project_id}/pdf")
def export_pdf(project_id: str):

    data = export_project(project_id)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer)

    styles = getSampleStyleSheet()
    elements = []

    # --- TITLE ---
    elements.append(Paragraph("FOUNDRA INVESTOR REPORT", styles["Title"]))
    elements.append(Spacer(1, 20))

    # --- PHASES ---
    for phase in data["phases"]:
        elements.append(
            Paragraph(f"Phase {phase['phase_number']}", styles["Heading2"])
        )
        elements.append(
            Paragraph(str(phase["output"]), styles["BodyText"])
        )
        elements.append(Spacer(1, 12))

    # --- BOARDROOM ---
    elements.append(Paragraph("Boardroom Decisions", styles["Heading2"]))
    for b in data["boardroom"]:
        elements.append(
            Paragraph(f"[Phase {b['phase']}] {b['status']}", styles["BodyText"])
        )
        elements.append(
            Paragraph(str(b["output"]), styles["BodyText"])
        )
        elements.append(Spacer(1, 10))

    doc.build(elements)

    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename=foundra_report_{project_id}.pdf"
        },
    )


@router.get("/api/exports/{project_id}/phases/{phase_number}")
def export_phase(project_id: str, phase_number: int, version: str = "latest"):

    query = (
        supabase.table("phase_runs")
        .select("*")
        .eq("project_id", project_id)
        .eq("phase_number", phase_number)
    )

    if version == "latest":
        query = query.order("retry_number", desc=True).limit(1)

    elif version == "all":
        query = query.order("retry_number", desc=False)

    data = query.execute().data

    return data


@router.get("/api/exports/{project_id}/phases/{phase_number}/boardroom")
def export_boardroom_phase(project_id: str, phase_number: int):

    res = (
        supabase.table("boardroom_runs")
        .select("*")
        .eq("project_id", project_id)
        .eq("phase_number", phase_number)
        .order("created_at", desc=False)
        .execute()
    )

    return {
        "phase_number": phase_number,
        "logs": res.data
    }



@router.get("/api/exports/{project_id}/phases/{phase_number}/pdf")
def export_phase_pdf(project_id: str, phase_number: int):
    # Fetch data
    phase_list = export_phase(project_id, phase_number)
    boardroom = export_boardroom_phase(project_id, phase_number)

    if not phase_list or len(phase_list) == 0:
        raise HTTPException(status_code=404, detail="Phase data not found")

    # FIX: Get the first item from the list
    phase = phase_list[0] 

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer)

    styles = getSampleStyleSheet()
    # Create a custom style for the output text to handle wrapping
    styles["BodyText"].wordWrap = 'CJK' 

    elements = []

    # Title
    elements.append(Paragraph(f"PHASE {phase_number} REPORT", styles["Title"]))
    elements.append(Spacer(1, 20))

    # Phase output
    elements.append(Paragraph("Phase Output", styles["Heading2"]))
    
    # Extract output safely
    output_data = phase.get("api_output") or phase.get("raw_output") or "No output data available."
    
    # If output is a dict, pretty print it so it doesn't look like a mess in the PDF
    if isinstance(output_data, (dict, list)):
        output_text = json.dumps(output_data, indent=2)
    else:
        output_text = str(output_data)

    elements.append(Paragraph(output_text.replace("\n", "<br/>"), styles["BodyText"]))
    elements.append(Spacer(1, 20))

    # Boardroom logs
    elements.append(Paragraph("Boardroom Decisions", styles["Heading2"]))

    if not boardroom.get("logs"):
        elements.append(Paragraph("No boardroom decisions recorded for this phase.", styles["BodyText"]))
    else:
        for log in boardroom["logs"]:
            status_color = "green" if log['status'] == "completed" else "red"
            elements.append(
                Paragraph(f"<font color='{status_color}'>{log['status'].upper()}</font> (retry {log['retry_number']})", styles["BodyText"])
            )
            
            log_output = log.get("api_output") or log.get("raw_output") or ""
            if isinstance(log_output, (dict, list)):
                log_text = json.dumps(log_output, indent=1)
            else:
                log_text = str(log_output)
                
            elements.append(Paragraph(log_text.replace("\n", "<br/>"), styles["BodyText"]))
            elements.append(Spacer(1, 10))

    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename=phase_{phase_number}_report.pdf"
        },
    )

@router.get("/api/exports/{project_id}/zip")
def export_zip(project_id: str):

    buffer = io.BytesIO()

    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:

        for phase_number in range(1, 7):  # assuming 6 phases

            try:
                phase = export_phase(project_id, phase_number)
                boardroom = export_boardroom_phase(project_id, phase_number)

                # --- Phase JSON
                z.writestr(
                    f"phases/phase_{phase_number}.json",
                    json.dumps(phase, indent=2)
                )

                # --- Boardroom JSON
                z.writestr(
                    f"boardroom/phase_{phase_number}_boardroom.json",
                    json.dumps(boardroom, indent=2)
                )

            except:
                continue

    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/zip",
        headers={
            "Content-Disposition": f"attachment; filename=foundra_export_{project_id}.zip"
        }
    )


#### MODERN STYLES FOR PDF EXPORT ####

def get_modern_styles():
    s = getSampleStyleSheet()
    # Professional Header Style
    s["Title"].fontName = "Helvetica-Bold"
    s["Title"].fontSize = 24
    s["Title"].spaceAfter = 30
    s["Title"].textColor = colors.HexColor("#1A1A1A")
    
    # Clean Body Style
    s["BodyText"].fontName = "Helvetica"
    s["BodyText"].fontSize = 10
    s["BodyText"].leading = 14
    s["BodyText"].alignment = 4  # Justified
    
    # Highlight Box for AI Insights
    s.add(getSampleStyleSheet()["BodyText"].__class__(
        name="InsightBox",
        parent=s["BodyText"],
        leftIndent=10,
        borderPadding=10,
        backColor=colors.HexColor("#F4F4F9")
    ))
    return s

@router.get("/api/exports/{project_id}/investor-pdf")
def investor_pdf(project_id: str):
    data = export_project(project_id)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=72)
    styles = get_modern_styles()
    elements = []

    # --- BRANDING & COVER ---
    elements.append(Paragraph("FOUNDRA", styles["Title"]))
    elements.append(Paragraph("VENTURE STRATEGY REPORT", styles["Heading2"]))
    elements.append(Paragraph(f"Generated on: {datetime.now().strftime('%B %d, %Y')}", styles["Italic"]))
    elements.append(Spacer(1, 0.5 * inch))

    # --- EXECUTIVE SUMMARY TABLE ---
    summary_data = [["PHASE", "STRATEGIC DOMAIN", "STATUS"]]
    domains = ["Problem Identification", "Market Validation", "Market Size", "Revenue Model", "MVP Architecture", "GTM Strategy"]
    
    for i, p in enumerate(data["phases"]):
        summary_data.append([f"Phase {p['phase_number']}", domains[i] if i < len(domains) else "Analysis", "✓ Verified"])

    t = Table(summary_data, colWidths=[1*inch, 3*inch, 1.5*inch])
    t.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 12),
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#eeeeee")),
        ('LINEBELOW', (0,0), (-1,0), 1, colors.black),
        ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 0.4 * inch))

    # --- PHASE DETAILS WITH SMART PARSING ---
    for p in data["phases"]:
        phase_title = domains[p['phase_number']-1] if p['phase_number'] <= len(domains) else f"Phase {p['phase_number']}"
        
        # Section Grouping
        section = [
            Paragraph(phase_title.upper(), styles["Heading2"]),
            Spacer(1, 6)
        ]
        
        output = p["output"]
        if isinstance(output, dict):
            for key, val in output.items():
                key_text = key.replace("_", " ").title()
                section.append(Paragraph(f"<b>{key_text}:</b> {val}", styles["BodyText"]))
                section.append(Spacer(1, 4))
        else:
            section.append(Paragraph(str(output), styles["BodyText"]))
            
        elements.append(KeepTogether(section))
        elements.append(Spacer(1, 15))

    doc.build(elements)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf")


##### PPTX EXPORTS ######

def add_slide(prs, title, content):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)

    slide.shapes.title.text = title
    slide.placeholders[1].text = content


from pptx.enum.text import PP_ALIGN

def add_modern_slide(prs, title_str, content_obj):
    # Use Layout 1 (Title and Content)
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    
    # Title Styling
    title = slide.shapes.title
    title.text = title_str
    
    # Content Body
    tf = slide.placeholders[1].text_frame
    tf.word_wrap = True
    
    if isinstance(content_obj, dict):
        for key, value in content_obj.items():
            p = tf.add_paragraph()
            p.text = f"{key.replace('_', ' ').title()}: {value}"
            p.level = 0
    else:
        tf.text = str(content_obj)

@router.get("/api/exports/{project_id}/pitch-deck")
def generate_pitch_deck(project_id: str):
    data = export_project(project_id)
    prs = Presentation()
    
    # 1. Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Foundra Venture Deck"
    slide.placeholders[1].text = f"Project ID: {project_id}\nAutomated Strategic Analysis"

    # 2. Map Phases to Slide Titles
    phase_titles = {
        1: "The Problem Space",
        2: "Validation & Feedback",
        3: "Market Opportunity (TAM/SAM)",
        4: "Business & Revenue Model",
        5: "The Solution (MVP)",
        6: "Execution & GTM"
    }

    for p in data["phases"]:
        title = phase_titles.get(p["phase_number"], f"Phase {p['phase_number']}")
        add_modern_slide(prs, title, p["output"])

    # 3. Vision Slide
    add_modern_slide(prs, "Future Vision", "Building the future with autonomous intelligence systems.")

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f"attachment; filename=foundra_pitch_{project_id}.pptx"}
    )

    data = export_project(project_id)

    prs = Presentation()

    # --- COVER ---
    add_slide(prs, "FOUNDRA", "AI-Powered Venture Intelligence")

    # --- PHASE MAPPING ---
    phase_map = {p["phase_number"]: p["output"] for p in data["phases"]}

    add_slide(prs, "Problem", str(phase_map.get(1, "N/A")))
    add_slide(prs, "Validation", str(phase_map.get(2, "N/A")))
    add_slide(prs, "Market Opportunity", str(phase_map.get(3, "N/A")))
    add_slide(prs, "Business Model", str(phase_map.get(4, "N/A")))
    add_slide(prs, "Product (MVP)", str(phase_map.get(5, "N/A")))
    add_slide(prs, "Go-To-Market", str(phase_map.get(6, "N/A")))

    # --- FINAL SLIDE ---
    add_slide(prs, "Vision", "Building the future with autonomous intelligence systems.")

    # --- SAVE ---
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={
            "Content-Disposition": f"attachment; filename=foundra_pitch_{project_id}.pptx"
        },
    )